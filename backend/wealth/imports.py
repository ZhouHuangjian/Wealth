import csv
import io
import hashlib
import uuid
import zipfile
from pathlib import Path
from decimal import Decimal
from django.conf import settings
from django.db import transaction
from .models import ImportBatch, SourceRecord, SourceFact, EvidenceLink, Account, Event
from .common import DomainError, get_obj, dec, day, digest, record, audit
from .ledger import post_event, reverse_event

ADAPTERS = [
    {
        "id": "standard_fund",
        "name": "标准基金确认记录 CSV / XLSX",
        "status": "synthetic_verified",
        "version": "fund-confirmation-1",
        "limitations": "按机构账单字段映射核对已有申购；并非广发、易方达等机构原生格式适配",
    },
    {
        "id": "generic",
        "name": "通用 CSV / XLSX 映射",
        "status": "synthetic_verified",
        "version": "generic-1",
        "limitations": "经合成文件验证；机构文件仍需字段映射与对账",
    },
    *[
        {
            "id": i,
            "name": n,
            "status": "planned",
            "version": None,
            "limitations": "缺少真实脱敏样本；可使用通用映射，不代表原生兼容",
        }
        for i, n in [
            ("miaomiao", "喵喵记账"),
            ("guangfa_futures", "广发期货"),
            ("yinhe_futures", "银河期货"),
            ("guangfa_funds", "广发基金"),
            ("efunds", "易方达基金"),
            ("alipay", "支付宝"),
            ("wechat", "微信"),
            ("bank", "银行"),
            ("broker", "港美券商"),
        ]
    ],
]

ALIASES = {
    "date": ["economic_date", "date", "交易时间", "交易日期", "时间", "日期"],
    "amount": ["amount", "金额", "金额(元)", "金额（元）", "交易金额"],
    "currency": ["currency", "币种"],
    "description": ["description", "备注", "商品", "商品名称", "交易对方", "摘要"],
    "external_id": ["external_id", "source_record_id", "交易单号", "订单号", "流水号"],
    "type": ["kind", "type", "收支", "收/支", "类型", "交易类型"],
    "category": ["category", "分类"],
}
TYPES = {
    "支出": "expense",
    "收入": "income",
    "退款": "refund",
    "转账": "transfer",
    "不计收支": "unclassified",
}


def parse_file(content, filename):
    if len(content) > 10 * 1024 * 1024:
        raise DomainError("文件超过 10MB 限制")
    ext = Path(filename).suffix.lower()
    if ext == ".csv" or ext == ".txt":
        decoded = None
        for enc in ("utf-8-sig", "gb18030"):
            try:
                decoded = content.decode(enc)
                break
            except UnicodeDecodeError:
                pass
        if decoded is None:
            raise DomainError("无法识别文本编码，请导出 UTF-8 CSV")
        if "\x00" in decoded:
            raise DomainError("文件不是有效文本")
        lines = decoded.splitlines()
        # Institutions may prefix explanatory lines; only recognize a credible header.
        start = 0
        for idx, line in enumerate(lines[:50]):
            if any(x in line for x in ALIASES["date"]) and any(
                x in line for x in ALIASES["amount"]
            ):
                start = idx
                break
        sample = "\n".join(lines[start:])
        try:
            dialect = csv.Sniffer().sniff(sample[:8192], delimiters=",\t;")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(sample), dialect=dialect)
        if not reader.fieldnames or len(reader.fieldnames) > 100:
            raise DomainError("无法识别有效表头")
        headers = [str(x).strip() for x in reader.fieldnames]
        if len(set(headers)) != len(headers):
            raise DomainError("存在重复列名，请先区分字段")
        rows = []
        for idx, row in enumerate(reader, start + 2):
            if None in row:
                raise DomainError(f"第 {idx} 行列数超过表头")
            clean = {str(k).strip(): str(v or "").strip() for k, v in row.items()}
            if any(clean.values()):
                rows.append((idx, clean))
            if len(rows) > 50000:
                raise DomainError("单文件最多 50,000 行，请拆分导出")
        return headers, rows
    if ext == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                if (
                    sum(i.file_size for i in z.infolist()) > 50 * 1024 * 1024
                    or len(z.infolist()) > 500
                ):
                    raise DomainError("表格解压后过大")
                if any(
                    i.file_size > max(i.compress_size, 1) * 200 for i in z.infolist()
                ):
                    raise DomainError("表格压缩比例异常")
            from openpyxl import load_workbook

            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
            sheet = book.active
            if sheet.max_row > 50001 or sheet.max_column > 100:
                raise DomainError("表格超过行列限制")
            values = sheet.iter_rows()
            headers = [str(c.value or "").strip() for c in next(values)]
            if len(set(headers)) != len(headers) or not all(headers):
                raise DomainError("表头为空或重复")
            rows = []
            for num, row in enumerate(values, 2):
                if any(c.data_type == "f" for c in row):
                    raise DomainError(f"第 {num} 行含公式，请导出固定值再导入")
                data = {
                    k: str(c.value if c.value is not None else "")
                    for k, c in zip(headers, row)
                }
                if any(data.values()):
                    rows.append((num, data))
            book.close()
            return headers, rows
        except DomainError:
            raise
        except Exception:
            raise DomainError("无法读取 XLSX，请检查文件是否损坏或加密")
    raise DomainError("当前支持 CSV、TXT、XLSX；PDF 与扫描件需专门适配")


def create_batch(space, user, uploaded, source, account_id=None):
    content = uploaded.read(10 * 1024 * 1024 + 1)
    checksum = hashlib.sha256(content).hexdigest()
    existing = ImportBatch.objects.filter(
        tenant=space, digest=checksum, source=source
    ).first()
    if existing:
        return existing
    headers, rows = parse_file(content, uploaded.name)
    if not rows:
        raise DomainError("文件没有可导入行")
    account = get_obj(Account, space, account_id) if account_id else None
    key = f"{space.pk}/{uuid.uuid4()}.source"
    target = settings.PRIVATE_MEDIA_ROOT / key
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_bytes(content)
    target.chmod(0o600)
    batch = ImportBatch.objects.create(
        tenant=space,
        created_by=user,
        source=source,
        filename=Path(uploaded.name.replace("\\", "/")).name[:200],
        digest=checksum,
        storage_key=key,
        account=account,
        diagnostics={
            "headers": headers,
            "row_count": len(rows),
            "compatibility": "generic_mapping_only",
        },
    )
    SourceRecord.objects.bulk_create(
        [
            SourceRecord(
                tenant=space, created_by=user, batch=batch, row_number=n, raw=r
            )
            for n, r in rows
        ],
        batch_size=1000,
    )
    audit(
        space,
        user,
        "import.uploaded",
        batch,
        {"row_count": len(rows), "digest": checksum},
    )
    return batch


@transaction.atomic
def preview(space, user, batch, body):
    if batch.status in {"committed", "reversed"}:
        raise DomainError("已提交批次不能原地重新解析", "conflict", 409)
    headers = batch.diagnostics["headers"]
    mapping = (
        body.get("mapping")
        or batch.mapping
        or {
            k: next((x for x in options if x in headers), "")
            for k, options in ALIASES.items()
        }
    )
    account = get_obj(Account, space, body.get("account_id") or batch.account_id)
    if not mapping.get("amount") or not mapping.get("date"):
        raise DomainError("必须映射日期和金额列")
    invalid = [v for v in mapping.values() if v and v not in headers]
    if invalid:
        raise DomainError("映射引用了不存在的列")
    committed_ids = EvidenceLink.objects.filter(
        tenant=space, record__batch=batch, active=True
    ).values_list("record_id", flat=True)
    rows = list(batch.rows.exclude(pk__in=committed_ids).order_by("row_number"))
    for row in rows:
        row.errors = []

        def value(field, default=""):
            return row.raw.get(mapping.get(field, ""), default)

        try:
            amt = dec(value("amount").replace(",", "").replace("¥", "").strip())
            raw_date = value("date").replace("/", "-").split(" ")[0]
            kind = (
                TYPES.get(value("type"), value("type"))
                or body.get("default_kind")
                or ("expense" if amt < 0 else "income")
            )
            if kind not in {
                "expense",
                "income",
                "transfer",
                "refund",
                "unclassified",
                "fund_debit",
                "dividend",
            }:
                raise DomainError("未知业务类型，请设置类型映射或单独处理")
            currency = value("currency") or body.get("currency") or account.currency
            if currency not in {"CNY", "HKD", "USD"} or currency != account.currency:
                raise DomainError("币种缺失或与账户不符")
            row.normalized = {
                "kind": kind,
                "account_id": str(account.pk),
                "economic_date": str(day(raw_date)),
                "amount": str(abs(amt)),
                "currency": currency,
                "description": value("description"),
                "category": value("category"),
                "external_id": value("external_id"),
            }
            if kind == "unclassified":
                direction = body.get("unclassified_direction")
                if direction in {"in", "out"}:
                    row.normalized["direction"] = direction
                elif amt < 0:
                    row.normalized["direction"] = "out"
                else:
                    row.errors.append(
                        "未分类流水方向未知，请核对后关联已有实际事项，不能推测扣款"
                    )
            if kind in {"transfer", "refund"}:
                row.errors.append("需映射目标账户或原消费；可先关联已有事项")
            native = value("external_id")
            row.identity = (
                digest([space.pk, batch.source, account.pk, native]) if native else ""
            )
        except (DomainError, ValueError) as e:
            row.normalized = {}
            row.errors.append(str(e))
            row.identity = ""
    SourceRecord.objects.bulk_update(
        rows, ["normalized", "errors", "identity"], batch_size=1000
    )
    batch.mapping = mapping
    batch.account = account
    batch.version += 1
    batch.ledger_revision = space.revision
    batch.status = "needs_review" if any(r.errors for r in rows) else "ready"
    batch.preview_hash = digest([[str(r.pk), r.normalized, r.errors] for r in rows])
    batch.save()
    return batch_preview(space, batch)


def batch_preview(space, batch, offset=0, limit=10000):
    facts = {
        f.identity: f
        for f in SourceFact.objects.filter(
            tenant=space,
            identity__in=batch.rows.exclude(identity="").values("identity"),
        )
    }
    candidates = default_candidates(space, batch)
    committed_ids = set(
        EvidenceLink.objects.filter(
            tenant=space, record__batch=batch, active=True
        ).values_list("record_id", flat=True)
    )
    result = record(batch)
    result["rows"] = []
    for r in batch.rows.all().order_by("row_number")[offset : offset + limit]:
        rr = record(r)
        f = facts.get(r.identity)
        rr["committed"] = r.pk in committed_ids
        rr["duplicate"] = bool(f)
        rr["existing_event_id"] = str(f.event_id) if f else None
        rr["candidates"] = candidates.get(
            (
                r.normalized.get("economic_date"),
                dec(r.normalized["amount"]) if r.normalized.get("amount") else None,
                r.normalized.get("currency"),
            ),
            [],
        )
        rr["automatic_match"] = any(
            c.get("automatic_estimate") for c in rr["candidates"]
        )
        result["rows"].append(rr)
    result["row_count"] = batch.rows.count()
    result["preview_version"] = batch.version
    return result


def default_candidates(space, batch):
    candidates = {}
    dates = [
        x for x in batch.rows.values_list("normalized__economic_date", flat=True) if x
    ]
    if not dates:
        return candidates
    for e in Event.objects.filter(
        tenant=space,
        payload__account_id=str(batch.account_id),
        economic_date__gte=min(dates),
        economic_date__lte=max(dates),
        reversal__isnull=True,
    ).exclude(kind="reversal")[:10000]:
        k = (
            str(e.economic_date),
            dec(e.payload.get("amount", "0")),
            e.payload.get("currency"),
        )
        candidates.setdefault(k, []).append(
            {
                "id": str(e.pk),
                "description": e.description,
                "reason": "同日同币种同金额候选；须核实资金账户与业务身份",
                "automatic_estimate": bool(e.payload.get("automatic_estimate")),
            }
        )
    return candidates


@transaction.atomic
def commit_batch(space, user, batch, body):
    if batch.status == "committed":
        return record(batch)
    if (
        int(body.get("preview_version", -1)) != batch.version
        or int(body.get("ledger_revision", -1)) != space.revision
        or batch.ledger_revision != space.revision
    ):
        raise DomainError("预览已过期，请重新预览后确认", "stale_preview", 409)
    if batch.status not in {"ready", "needs_review"}:
        raise DomainError("请先完成映射预览")
    ids = body.get("row_ids")
    links = body.get("links", {})
    distinct = body.get("distinct_rows", [])
    if not isinstance(distinct, list) or any(not isinstance(v, str) for v in distinct):
        raise DomainError("独立交易确认格式不正确")
    rows = list(batch.rows.filter(pk__in=ids) if ids is not None else batch.rows.all())
    if ids is not None and len(rows) != len(set(ids)):
        raise DomainError("提交行不属于本批次")
    if not set(distinct).issubset({str(r.pk) for r in rows}):
        raise DomainError("独立交易确认不属于提交行")
    linked = set(
        EvidenceLink.objects.filter(
            tenant=space, record__batch=batch, active=True
        ).values_list("record_id", flat=True)
    )
    rows = [row for row in rows if row.pk not in linked]
    if not rows:
        raise DomainError("没有选中可提交行")
    # Validate the entire selected set before publishing; outer transaction rolls back all effects.
    for row in rows:
        if row.errors and str(row.pk) not in links:
            raise DomainError(f"第 {row.row_number} 行仍有错误")
    for row in rows:
        normalized = row.normalized
        f = (
            SourceFact.objects.filter(tenant=space, identity=row.identity)
            .select_related("event")
            .first()
            if row.identity
            else None
        )
        introduced = False
        if str(row.pk) in links:
            event = get_obj(Event, space, links[str(row.pk)])
            if not normalized:
                raise DomainError("请先修正来源字段，再关联已有事项")
            if event.payload.get("currency") != normalized.get("currency") or dec(
                event.payload.get("amount", "0")
            ) != dec(normalized["amount"]):
                raise DomainError("关联金额或币种不符，需解释费用或拆分后处理")
            if (
                event.payload.get("account_id") != normalized["account_id"]
                or str(event.economic_date) != normalized["economic_date"]
                or (
                    normalized["kind"] != "unclassified"
                    and event.kind != normalized["kind"]
                )
            ):
                raise DomainError("关联的账户、日期或业务类型不符，请核实来源身份")
            if f and f.event_id != event.pk:
                raise DomainError(
                    "来源身份已关联另一事项，不能覆盖", "source_revision", 409
                )
            if hasattr(event, "reversal"):
                raise DomainError("不能关联已冲正事项")
        elif f:
            if f.fingerprint != digest(normalized):
                raise DomainError(
                    "同来源记录内容已修订，须显式更正而非重复覆盖",
                    "source_revision",
                    409,
                )
            event = f.event
            if hasattr(event, "reversal"):
                raise DomainError(
                    "原来源事项已冲正，需建立明确的替代链", "source_revision", 409
                )
        else:
            if normalized.get("kind") == "fund_debit" and str(row.pk) not in distinct:
                automatic = Event.objects.filter(
                    tenant=space,
                    kind="fund_debit",
                    reversal__isnull=True,
                    economic_date=normalized["economic_date"],
                    payload__account_id=normalized["account_id"],
                    payload__currency=normalized["currency"],
                    payload__automatic_estimate=True,
                )
                if any(
                    dec(e.payload.get("amount", "0")) == dec(normalized["amount"])
                    for e in automatic
                ):
                    raise DomainError(
                        f"第 {row.row_number} 行已有同金额的自动定投记录，请关联已有事项，或明确选择这是另一笔独立交易",
                        "automatic_duplicate",
                        409,
                    )
            event = post_event(
                space,
                user,
                normalized,
                stage_key=f"{space.pk}:source:{row.identity or row.pk}",
            )
            introduced = True
        if row.identity and not f:
            SourceFact.objects.create(
                tenant=space,
                created_by=user,
                identity=row.identity,
                event=event,
                fingerprint=digest(normalized),
            )
        EvidenceLink.objects.get_or_create(
            tenant=space,
            record=row,
            event=event,
            defaults={"created_by": user, "introduced": introduced},
        )
    committed_count = (
        EvidenceLink.objects.filter(tenant=space, record__batch=batch, active=True)
        .values("record_id")
        .distinct()
        .count()
    )
    remaining = batch.rows.count() - committed_count
    batch.status = "partially_committed" if remaining else "committed"
    batch.diagnostics.update(
        committed_rows=committed_count, remaining_rows=remaining, excluded_rows=0
    )
    batch.save()
    audit(space, user, "import.committed", batch, {"count": len(rows)})
    return record(batch)


@transaction.atomic
def reverse_batch(space, user, batch, reason):
    if batch.parser_version == "fund-confirmation-1":
        from .fund_reconciliation import reverse

        return reverse(space, user, batch, reason)
    if batch.status not in {"committed", "partially_committed"}:
        raise DomainError("只能撤销含已提交事项的批次")
    if not reason:
        raise DomainError("请填写撤销原因")
    links = list(
        EvidenceLink.objects.filter(tenant=space, record__batch=batch, active=True)
        .select_related("event")
        .order_by("-event__created_at")
    )
    for link in links:
        other = (
            EvidenceLink.objects.filter(tenant=space, event=link.event, active=True)
            .exclude(record__batch=batch)
            .exists()
        )
        if link.introduced and not other and not hasattr(link.event, "reversal"):
            reverse_event(space, user, link.event, reason)
        link.active = False
        link.save(update_fields=["active"])
    batch.status = "reversed"
    batch.save()
    audit(space, user, "import.reversed", batch, {"reason": reason})
    return record(batch)
